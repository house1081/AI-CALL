#!/usr/bin/env python3
"""导出主线+FAQ 文本/录音对照表为 CSV + XLSX。"""
from __future__ import annotations

import csv
import json
import re
import urllib.request
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
PACK = json.loads(
    (ROOT / "ai-call-server/src/main/resources/dialog-scripts/loan-pack.json").read_text(encoding="utf-8")
)
UPLOAD_LOG_PATH = ROOT / "scripts/_upload_result_latest.txt"
if not UPLOAD_LOG_PATH.exists():
    UPLOAD_LOG_PATH = ROOT / "scripts/_upload_result.txt"
def read_upload_log(path: Path) -> str:
    for enc in ("utf-8", "utf-8-sig", "utf-16", "gbk"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    return path.read_bytes().decode("utf-8", errors="replace")


UPLOAD_LOG = read_upload_log(UPLOAD_LOG_PATH) if UPLOAD_LOG_PATH.exists() else ""
AUDIO_DIR = ROOT / "ai-call-server" / "录音文件"
OUT_CSV = ROOT / "scripts" / "kb_audio对照表.csv"
OUT_XLSX = ROOT / "scripts" / "kb_audio对照表.xlsx"
BASE = "http://127.0.0.1:8081/api/admin"
FIELDNAMES = ["类型", "编号", "场景/关键词", "标准话术", "QA_ID", "已有WAV", "WAV路径", "本地m4a", "建议操作"]


def login() -> str:
    body = json.dumps({"username": "admin", "password": "admin123"}).encode()
    req = urllib.request.Request(
        f"{BASE}/login", data=body, headers={"Content-Type": "application/json"}
    )
    with urllib.request.urlopen(req, timeout=15) as resp:
        data = json.loads(resp.read())
    if data.get("code") != 200:
        raise RuntimeError(data.get("message") or "login failed")
    return data["data"]["token"]


def fetch_qas(token: str) -> list[dict]:
    req = urllib.request.Request(
        f"{BASE}/dialog-training/list?kbId=1&status=1",
        headers={"Authorization": f"Bearer {token}"},
    )
    with urllib.request.urlopen(req, timeout=60) as resp:
        data = json.loads(resp.read())
    if data.get("code") != 200:
        raise RuntimeError(data.get("message") or "list failed")
    return data["data"]


def wav_ok(path: str | None) -> bool:
    if not path:
        return False
    p = Path(path.replace("\\", "/").replace("./", ""))
    if not p.is_absolute():
        p = ROOT / "ai-call-server" / p
    return p.exists() and p.stat().st_size > 44


def parse_upload_log() -> dict[str, list[str]]:
    flow: dict[str, list[str]] = {}
    fb: dict[str, list[str]] = {}
    for line in UPLOAD_LOG.splitlines():
        m = re.match(r"\[上传\] (.+?) -> qa:(\d+) \((flow|fallback):([^)]+)\)", line)
        if m:
            stem, typ, key = m.group(1), m.group(3), m.group(4)
            (flow if typ == "flow" else fb).setdefault(key, []).append(stem)
    return {"flow": flow, "fallback": fb}


def suggest_action(typ: str, key: str, has_wav: bool, m4a: list[str]) -> str:
    if has_wav:
        return "已完成"
    if m4a:
        return "上传m4a"
    return "需新录"


def write_xlsx(rows: list[dict]) -> None:
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1E3A5F")
    wrap = Alignment(wrap_text=True, vertical="top")

    def fill_sheet(ws, data: list[dict], title: str) -> None:
        ws.title = title[:31]
        ws.append(FIELDNAMES)
        for col in range(1, len(FIELDNAMES) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in data:
            ws.append([row.get(k, "") for k in FIELDNAMES])
        for r in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in r:
                cell.alignment = wrap
                if cell.column == 6 and cell.value == "否":
                    cell.fill = PatternFill("solid", fgColor="FDE2E2")
                if cell.column == 9 and cell.value and "新录" in str(cell.value):
                    cell.fill = PatternFill("solid", fgColor="FFF3CD")
                if cell.column == 9 and cell.value == "已完成":
                    cell.fill = PatternFill("solid", fgColor="E1F3D8")
        widths = [8, 8, 28, 52, 10, 10, 36, 24, 12]
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    fill_sheet(wb.active, rows, "全部对照")
    need = [r for r in rows if "新录" in r.get("建议操作", "")]
    ws2 = wb.create_sheet("需新录")
    fill_sheet(ws2, need, "需新录")
    done = [r for r in rows if r.get("建议操作") == "已完成"]
    ws3 = wb.create_sheet("已完成")
    fill_sheet(ws3, done, "已完成")

    summary = wb.create_sheet("统计", 0)
    summary.append(["项目", "数量"])
    summary["A1"].font = header_font
    summary["B1"].font = header_font
    summary["A1"].fill = header_fill
    summary["B1"].fill = header_fill
    total = len(rows)
    done_n = len(done)
    need_n = len(need)
    main_need = sum(1 for r in need if r["类型"] == "主线")
    faq_need = sum(1 for r in need if r["类型"] == "FAQ")
    for label, val in [
        ("总行数", total),
        ("已有WAV（已完成）", done_n),
        ("需新录合计", need_n),
        ("  其中主线", main_need),
        ("  其中FAQ", faq_need),
        ("导出时间数据源", "实时API + loan-pack"),
    ]:
        summary.append([label, val])
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 16

    wb.save(OUT_XLSX)


def main() -> None:
    token = login()
    qas = fetch_qas(token)
    qa_by_remark = {(q.get("remark") or ""): q for q in qas}
    upload_map = parse_upload_log()
    m4a_all = sorted(p.stem for p in AUDIO_DIR.glob("*.m4a"))

    rows: list[dict] = []

    for step in sorted(PACK["mainFlow"], key=lambda x: (len(x["step"]), x["step"])):
        remark = f"flow:{step['step']}"
        qa = qa_by_remark.get(remark, {})
        m4a = upload_map["flow"].get(step["step"], [])
        has = wav_ok(qa.get("answerWavPath"))
        rows.append({
            "类型": "主线",
            "编号": step["step"],
            "场景/关键词": step["scene"],
            "标准话术": step["script"],
            "QA_ID": qa.get("id", ""),
            "已有WAV": "是" if has else "否",
            "WAV路径": qa.get("answerWavPath") or "",
            "本地m4a": " | ".join(m4a),
            "建议操作": suggest_action("flow", step["step"], has, m4a),
        })

    for fb in sorted(PACK["fallbacks"], key=lambda x: x["no"]):
        remark = f"fallback:{fb['no']}"
        qa = qa_by_remark.get(remark, {})
        m4a = upload_map["fallback"].get(str(fb["no"]), [])
        has = wav_ok(qa.get("answerWavPath"))
        rows.append({
            "类型": "FAQ",
            "编号": fb["no"],
            "场景/关键词": fb["keywords"],
            "标准话术": fb["answer"],
            "QA_ID": qa.get("id", ""),
            "已有WAV": "是" if has else "否",
            "WAV路径": qa.get("answerWavPath") or "",
            "本地m4a": " | ".join(m4a),
            "建议操作": suggest_action("fallback", str(fb["no"]), has, m4a),
        })

    rows.append({
        "类型": "话术模板",
        "编号": "opening",
        "场景/关键词": "外呼开场白",
        "标准话术": PACK["mainFlow"][0]["script"],
        "QA_ID": "",
        "已有WAV": "",
        "WAV路径": "",
        "本地m4a": "开场白" if "开场白" in m4a_all else "",
        "建议操作": "上传至AI话术开场白",
    })

    fieldnames = FIELDNAMES
    with OUT_CSV.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(rows)

    write_xlsx(rows)

    done = sum(1 for r in rows if r["已有WAV"] == "是")
    upload = sum(1 for r in rows if r["建议操作"] == "上传m4a")
    record = sum(1 for r in rows if "新录" in r["建议操作"])
    print(f"已导出 {len(rows)} 行")
    print(f"  CSV  -> {OUT_CSV}")
    print(f"  XLSX -> {OUT_XLSX}")
    print(f"已有WAV: {done} | 待上传m4a: {upload} | 需新录: {record}")


if __name__ == "__main__":
    main()
