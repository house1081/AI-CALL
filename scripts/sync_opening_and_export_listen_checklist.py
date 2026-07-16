#!/usr/bin/env python3
"""1) 统一 ai-prompt openingRemarks 与 loan-pack 01，并上传开场白录音
   2) 导出多段 FAQ + 语义风险听音核对清单 Excel"""
from __future__ import annotations

import json
import mimetypes
import re
import urllib.error
import urllib.request
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
PACK_PATH = ROOT / "ai-call-server/src/main/resources/dialog-scripts/loan-pack.json"
OPENING_M4A = ROOT / "ai-call-server/录音文件-6-23/开场白.m4a"
if not OPENING_M4A.exists():
    OPENING_M4A = ROOT / "ai-call-server/录音文件/开场白.m4a"
OUT_XLSX = ROOT / "scripts/听音核对清单.xlsx"
BASE = "http://127.0.0.1:8081/api/admin"

SEMANTIC_RISKS = [
    ("FAQ", "19", "听不清|重复一遍", 227,
     "听不清.m4a",
     "标准答是重播开场白全文，不是「听不清」类应答；需听音确认 m4a 内容"),
    ("FAQ", "21", "信号不好", 229,
     "信号不好.m4a",
     "标准答「喂，您好，您这边能听见么？」；勿与 FAQ42「在听的，您接着说」混用"),
    ("FAQ", "42", "听得到吗|听得见吗", 250,
     "听得见吗.m4a / 喂？.m4a",
     "标准答「在听的，您接着说」；确认 m4a 与文本一致"),
    ("FAQ", "101", "静默/没声音", 309,
     "喂？.m4a / 喂？在听吗？.m4a",
     "标准答 3 段（喂？/在听吗/再见），单 wav 可能只录一句"),
    ("FAQ", "97", "没听清|表述模糊", 305,
     "你说什么.m4a 等",
     "标准答 3 段分号分隔，录音可能不完整"),
]


def login() -> str:
    body = json.dumps({"username": "admin", "password": "admin123"}).encode()
    req = urllib.request.Request(
        f"{BASE}/login", data=body, headers={"Content-Type": "application/json"}
    )
    with urllib.request.urlopen(req, timeout=15) as resp:
        data = json.loads(resp.read())
    if data.get("code") != 200:
        raise RuntimeError(data.get("message") or "login failed")
    return data["data"]["token"]


def api_get(token: str, path: str):
    req = urllib.request.Request(f"{BASE}{path}", headers={"Authorization": f"Bearer {token}"})
    with urllib.request.urlopen(req, timeout=60) as resp:
        data = json.loads(resp.read())
    if data.get("code") != 200:
        raise RuntimeError(data.get("message") or f"GET {path}")
    return data["data"]


def api_post_json(token: str, path: str, payload: dict):
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    req = urllib.request.Request(
        f"{BASE}{path}",
        data=body,
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type": "application/json; charset=utf-8",
        },
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=60) as resp:
        data = json.loads(resp.read())
    if data.get("code") != 200:
        raise RuntimeError(data.get("message") or f"POST {path}")
    return data.get("data")


def upload_opening(token: str, prompt_id: int, file_path: Path):
    boundary = "----AiCallOpeningSync7MA4YWxk"
    mime = mimetypes.guess_type(file_path.name)[0] or "application/octet-stream"
    content = file_path.read_bytes()
    body = (
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="file"; filename="{file_path.name}"\r\n'
        f"Content-Type: {mime}\r\n\r\n"
    ).encode("utf-8") + content + f"\r\n--{boundary}--\r\n".encode()
    req = urllib.request.Request(
        f"{BASE}/ai-prompt/{prompt_id}/upload-opening-audio",
        data=body,
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type": f"multipart/form-data; boundary={boundary}",
        },
        method="POST",
    )
    with urllib.request.urlopen(req, timeout=120) as resp:
        data = json.loads(resp.read())
    if data.get("code") != 200:
        raise RuntimeError(data.get("message") or "upload opening failed")
    return data["data"]


def sync_opening(token: str) -> None:
    pack = json.loads(PACK_PATH.read_text(encoding="utf-8"))
    opening_text = next(x["script"] for x in pack["mainFlow"] if x["step"] == "01")

    prompts = api_get(token, "/ai-prompt/list")
    active = next(p for p in prompts if p.get("isActive") == 1)
    pid = active["id"]
    old = (active.get("openingRemarks") or "").strip()

    print("=== 统一外呼开场白 ===")
    print(f"  prompt id={pid}")
    print(f"  旧 openingRemarks: {old[:60]}…" if len(old) > 60 else f"  旧 openingRemarks: {old}")
    print(f"  新 openingRemarks: {opening_text[:60]}…" if len(opening_text) > 60 else f"  新 openingRemarks: {opening_text}")

    if old != opening_text.strip():
        active["openingRemarks"] = opening_text
        api_post_json(token, "/ai-prompt/save", active)
        print("  [OK] openingRemarks 已更新为 loan-pack 01")
    else:
        print("  [跳过] openingRemarks 已与 loan-pack 01 一致")

    if not OPENING_M4A.exists():
        print(f"  [警告] 未找到开场白 m4a: {OPENING_M4A}")
        return

    result = upload_opening(token, pid, OPENING_M4A)
    print(f"  [OK] 开场白录音已上传: {result.get('openingWavPath')}")
    print(f"       音频 URL: {result.get('audioUrl')}")


def fetch_multi_segment_and_risks(token: str) -> tuple[list[dict], list[dict]]:
    qas = api_get(token, "/dialog-training/list?kbId=1&status=1")
    pack = json.loads(PACK_PATH.read_text(encoding="utf-8"))
    fb_pack = {str(x["no"]): x for x in pack["fallbacks"]}

    multi: list[dict] = []
    for q in qas:
        remark = q.get("remark") or ""
        if not remark.startswith("fallback:"):
            continue
        ans = (q.get("standardAnswer") or "").strip()
        if "；" not in ans and ";" not in ans:
            continue
        no = remark.split(":", 1)[1]
        segs = [s.strip() for s in re.split("[；;]", ans) if s.strip()]
        pack_row = fb_pack.get(no, {})
        multi.append({
            "类型": "FAQ",
            "编号": no,
            "场景/关键词": pack_row.get("keywords", q.get("question", "")),
            "QA_ID": q["id"],
            "段数": len(segs),
            "第1段": segs[0] if segs else "",
            "第2段": segs[1] if len(segs) > 1 else "",
            "第3段及以后": "；".join(segs[2:]) if len(segs) > 2 else "",
            "完整标准答": ans,
            "已有WAV": "是" if q.get("answerWavPath") else "否",
            "WAV路径": q.get("answerWavPath") or "",
            "核对项": "听音：wav 是否只录了第1段？若否请重录完整或拆条",
            "核对结果": "",
            "备注": "",
        })

    risks: list[dict] = []
    qa_by_id = {q["id"]: q for q in qas}
    for typ, no, kw, qa_id, m4a_hint, note in SEMANTIC_RISKS:
        q = qa_by_id.get(qa_id, {})
        risks.append({
            "类型": typ,
            "编号": no,
            "场景/关键词": kw,
            "QA_ID": qa_id,
            "段数": "",
            "第1段": "",
            "第2段": "",
            "第3段及以后": "",
            "完整标准答": q.get("standardAnswer", ""),
            "已有WAV": "是" if q.get("answerWavPath") else "否",
            "WAV路径": q.get("answerWavPath") or "",
            "核对项": note,
            "本地m4a提示": m4a_hint,
            "核对结果": "",
            "备注": "",
        })
    return multi, risks


def write_excel(multi: list[dict], risks: list[dict]) -> None:
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "多段FAQ听音核对"
    ws2 = wb.create_sheet("语义风险听音核对")

    headers_multi = [
        "类型", "编号", "场景/关键词", "QA_ID", "段数",
        "第1段", "第2段", "第3段及以后", "完整标准答",
        "已有WAV", "WAV路径", "核对项", "核对结果", "备注",
    ]
    headers_risk = [
        "类型", "编号", "场景/关键词", "QA_ID",
        "完整标准答", "已有WAV", "WAV路径",
        "本地m4a提示", "核对项", "核对结果", "备注",
    ]

    def fill_sheet(ws, headers: list[str], rows: list[dict]):
        header_fill = PatternFill("solid", fgColor="4472C4")
        header_font = Font(color="FFFFFF", bold=True)
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for ri, row in enumerate(rows, 2):
            for ci, h in enumerate(headers, 1):
                ws.cell(row=ri, column=ci, value=row.get(h, ""))
                ws.cell(row=ri, column=ci).alignment = Alignment(vertical="top", wrap_text=True)
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18
        ws.column_dimensions["I"].width = 36
        ws.column_dimensions["J"].width = 28

    fill_sheet(ws1, headers_multi, multi)
    fill_sheet(ws2, headers_risk, risks)
    ws2.column_dimensions["E"].width = 40
    ws2.column_dimensions["H"].width = 24
    ws2.column_dimensions["I"].width = 42

    wb.save(OUT_XLSX)
    print(f"\n=== 听音核对清单 ===")
    print(f"  多段 FAQ: {len(multi)} 条")
    print(f"  语义风险: {len(risks)} 条")
    print(f"  已导出: {OUT_XLSX}")


def main() -> int:
    try:
        token = login()
    except urllib.error.URLError as e:
        print(f"无法连接后端 {BASE}: {e}")
        print("请先启动 ai-call-server 后再运行本脚本。")
        return 1

    sync_opening(token)
    multi, risks = fetch_multi_segment_and_risks(token)
    write_excel(multi, risks)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
